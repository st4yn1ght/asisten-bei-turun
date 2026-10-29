"""T-LP: laporan harian (SPEC §18.1–§18.2, §18.4) tanpa bagian Telegram (tahap 7).

Setiap keluaran memuat "BUKAN SARAN INVESTASI"; HTML mandiri (tanpa JavaScript/URL eksternal, gambar PNG base64);
instruksi pembulatan dan instruksi menarik order muncul; berkas terkunci → nama alternatif + peringatan.
Proyek sintetis `bantu_harian` (FIKTIF – KHUSUS UJI).
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from asisten_bei_turun import operasional
from asisten_bei_turun.laporan.tabel import SHEET

SEMUA_FORMAT = {"laporan.format": ["terminal", "markdown", "html", "csv", "xlsx"]}
DISCLAIMER = "BUKAN SARAN INVESTASI"
T = "2025-06-30"  # BOLEH BELI (K1); REKOMENDASI AAAA (TB) dan CCCC (PJ) pada data sintetis


@pytest.fixture
def lengkap(proyek):
    proyek.ubah_config(SEMUA_FORMAT)
    h = proyek.harian(T)
    assert h.kode == 0, h.teks
    return proyek, h


def test_semua_berkas_laporan_ditulis_dan_memuat_disclaimer(lengkap):
    p, h = lengkap
    folder = p.folder_laporan / T
    nama = {"laporan.md", "laporan.html", "kandidat.csv", "rencana.xlsx", "kualitas_data.csv"}
    assert nama <= {x.name for x in folder.iterdir()}
    run = [x for x in folder.iterdir() if x.is_dir()]
    assert len(run) == 1 and nama <= {x.name for x in run[0].iterdir()}
    assert h.teks.startswith(DISCLAIMER) and h.baris[-1].startswith(DISCLAIMER)
    for n in ("kandidat.csv", "kualitas_data.csv"):
        assert (folder / n).read_text(encoding="utf-8").startswith("# " + DISCLAIMER)
    md = (folder / "laporan.md").read_text(encoding="utf-8")
    assert md.startswith(DISCLAIMER) and md.rstrip().endswith("tanggung jawab pengguna.")
    wb = load_workbook(folder / "rencana.xlsx")
    assert wb.sheetnames == list(SHEET)
    for ws in wb.worksheets:
        assert str(ws["A1"].value).startswith(DISCLAIMER), ws.title


def test_semua_keluaran_teks_di_output_memuat_disclaimer(lengkap):
    p, _h = lengkap
    for f in (p.akar / "output").rglob("*"):
        if f.suffix in (".md", ".html", ".csv", ".json", ".yaml"):
            assert DISCLAIMER in f.read_text(encoding="utf-8"), f


def test_html_mandiri(lengkap):
    p, _h = lengkap
    html = (p.folder_laporan / T / "laporan.html").read_text(encoding="utf-8")
    assert "<script" not in html.lower()
    assert not re.search(r"(src|href)\s*=\s*[\"'](?!data:)", html)
    assert "http://" not in html and "https://" not in html
    gambar = re.findall(r'<img [^>]*src="data:image/png;base64,', html)
    assert len(gambar) == 4  # IHSG, breadth, AAAA dan CCCC (rencana)
    assert html.count(DISCLAIMER) >= 2 and "Glosarium" in html


def test_instruksi_pembulatan_dan_menarik_order(lengkap):
    p, h = lengkap
    md = (p.folder_laporan / T / "laporan.md").read_text(encoding="utf-8")
    assert "tarik (withdraw) order" in md and "tarik sisanya" in md
    assert "bulatkan limit beli dan target ke BAWAH, stop ke ATAS" in md
    assert "Pasang order otomatis stop di" in md
    assert "tarik jika tidak terisi saat pembukaan" in h.teks


def test_bagian_laporan_berurutan_spec_18_1(lengkap):
    p, _h = lengkap
    md = (p.folder_laporan / T / "laporan.md").read_text(encoding="utf-8")
    judul = [b[3:] for b in md.splitlines() if b.startswith("## ")]
    assert judul == ["Ringkasan aksi", "Status pasar", "Status strategi", "Akun", "Posisi dipegang", "Rencana beli",
                     "Kandidat lain", "Daftar pantau", "Kualitas data", "Item yang belum diverifikasi", "Jurnal kertas",
                     "Peringatan", "Penutup", "Glosarium"]


def test_berkas_terkunci_ditulis_ke_nama_alternatif(proyek, monkeypatch):
    proyek.ubah_config(SEMUA_FORMAT)
    asli = operasional.tulis_atomik
    terkunci = proyek.folder_laporan / T / "rencana.xlsx"

    def tulis(path, isi, encoding="utf-8"):
        if Path(path) == terkunci:
            raise PermissionError("terkunci (dibuka di Excel)")
        return asli(path, isi, encoding)

    monkeypatch.setattr(operasional, "tulis_atomik", tulis)
    h = proyek.harian(T)
    assert h.kode == 0
    alt = proyek.folder_laporan / T / "rencana_200000.xlsx"
    assert alt.is_file() and not terkunci.exists()
    assert "sedang terkunci" in h.teks
    md = (proyek.folder_laporan / T / "laporan.md").read_text(encoding="utf-8")
    assert "rencana_200000.xlsx" in md


def test_posisi_dipegang_tampil_dengan_level_order(proyek):
    proyek.tulis_input(
        "portfolio.csv",
        "kode,tanggal_beli,harga_beli,lot,strategi,stop_awal,sektor,harga_per_tanggal,lembar,catatan\n"
        "BBBB,2025-06-24,2400,1,TB,2200,,,,FIKTIF\n",
    )
    proyek.ubah_config(SEMUA_FORMAT)
    h = proyek.harian(T)
    assert h.kode == 0, h.teks
    md = (proyek.folder_laporan / T / "laporan.md").read_text(encoding="utf-8")
    bagian = md.split("## Posisi dipegang")[1].split("## ")[0]
    assert "| BBBB | TB |" in bagian and "Stop t+1 (order otomatis)" in bagian
    html = (proyek.folder_laporan / T / "laporan.html").read_text(encoding="utf-8")
    assert len(re.findall(r'<img [^>]*src="data:image/png;base64,', html)) == 5
    wb = load_workbook(proyek.folder_laporan / T / "rencana.xlsx")
    posisi = [r for r in wb["Posisi"].iter_rows(values_only=True)]
    assert posisi[2][0] == "BBBB"
