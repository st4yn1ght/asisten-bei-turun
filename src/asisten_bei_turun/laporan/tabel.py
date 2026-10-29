"""`kandidat.csv` dan `rencana.xlsx` (SPEC §18.2).

- `kandidat.csv`: semua sinyal hari itu beserta status (urutan seleksi §13.2); baris pertama disclaimer. Format
  mesin seperti berkas backtest: titik desimal tanpa pemisah ribuan, Rupiah 2 desimal, harga sampai 6 desimal.
- `rencana.xlsx`: sheet `Ringkasan`, `Rencana`, `Kandidat`, `Posisi`, `Pasar`, `Kualitas Data`, `Catatan`;
  baris 1 setiap sheet = disclaimer. Angka ditulis sebagai angka (Decimal), teks untuk dibaca pengguna.
"""

from __future__ import annotations

import io
from decimal import Decimal

import pandas as pd

from .. import DISCLAIMER
from ..backtest.laporan_bt import csv_harga, csv_rupiah, csv_tanggal
from ..jurnal import target_untuk
from ..aturan_bei import AturanBEI
from ..operasional import teks_csv
from .dokumen import _baris_kertas
from .model import KETERBATASAN_SINGKAT, LaporanHarian, rp, status_pasar_ringkas

KOLOM_KANDIDAT: tuple[str, ...] = (
    "tanggal_sinyal", "kode", "nama", "sektor", "strategi", "skor", "status_sinyal", "alasan", "status_pasar",
    "harga_ref", "harga_maks", "stop_awal", "target_harga", "lot", "lembar", "nilai_idr", "biaya_beli_idr",
    "risiko_idr", "rata_nilai_idr", "bergantung_jual",
)
SHEET: tuple[str, ...] = ("Ringkasan", "Rencana", "Kandidat", "Posisi", "Pasar", "Kualitas Data", "Catatan")


def baris_kandidat(lap: LaporanHarian) -> list[dict]:
    bei = AturanBEI(lap.cfg)
    hasil = []
    for k in lap.kandidat:
        sg = k.sinyal.sinyal
        sz = k.sizing
        rekom = k.status == "REKOMENDASI"
        hasil.append({
            "tanggal_sinyal": csv_tanggal(sg.tanggal), "kode": k.kode, "nama": lap.nama.get(k.kode, ""),
            "sektor": lap.sektor.get(k.kode, ""), "strategi": k.strategi, "skor": csv_harga(round(k.sinyal.skor, 6)),
            "status_sinyal": k.status, "alasan": " ".join((k.alasan or "").split()), "status_pasar": lap.status_pasar or "",
            "harga_ref": csv_harga(sg.harga_ref), "harga_maks": csv_harga(sg.harga_maks), "stop_awal": csv_harga(sg.stop),
            "target_harga": csv_harga(target_untuk(sg, bei)),
            "lot": str(sz.lot) if sz is not None else "", "lembar": str(sz.lembar) if sz is not None else "",
            "nilai_idr": csv_rupiah(sz.nilai_idr) if sz is not None and rekom else "",
            "biaya_beli_idr": csv_rupiah(sz.biaya_beli_idr) if sz is not None and rekom else "",
            "risiko_idr": csv_rupiah(sz.risiko_idr) if sz is not None and rekom else "",
            "rata_nilai_idr": csv_rupiah(Decimal(repr(float(sg.rata_nilai)))),
            "bergantung_jual": ";".join("+".join(g) for g in k.bergantung_jual),
        })
    return hasil


def teks_kandidat_csv(lap: LaporanHarian) -> str:
    return teks_csv(pd.DataFrame(baris_kandidat(lap), columns=list(KOLOM_KANDIDAT), dtype=object))


def _r(x: Decimal | None, desimal: int = 2):
    if x is None:
        return None
    return Decimal(x).quantize(Decimal(1).scaleb(-desimal))


def isi_sheet(lap: LaporanHarian) -> dict[str, list[list]]:
    """Isi setiap sheet sebagai baris nilai (baris pertama = disclaimer)."""
    d = [[DISCLAIMER]]
    ringkasan = d + [[x] for x in lap.kepala()] + [["BANNER: " + b] for b in lap.banner]
    ringkasan += [[], ["Ringkasan aksi"]] + [[a] for a in lap.ringkasan_aksi]
    ringkasan += [[], ["Akun", ""]] + [[a, b] for a, b in lap.akun_baris()]
    ringkasan += [[], ["Status pasar", status_pasar_ringkas(lap.kondisi)]]

    rencana = d + [["kode", "nama", "sektor", "strategi", "skor", "harga_ref", "harga_maks", "stop", "metode_stop",
                    "target", "lot", "lembar", "nilai_idr", "fee_beli_idr", "risiko_idr", "risiko_pct_ekuitas",
                    "rr_bersih", "instruksi_masuk", "ketergantungan", "langkah_setelah_terisi", "peringatan"]]
    for r in lap.rencana:
        rencana.append([r.kode, r.nama, r.sektor, r.strategi, Decimal(r.skor.replace(",", ".")), r.harga_ref,
                        r.harga_maks, r.stop, r.metode_stop, r.target[-1][2] if r.target else r.teks_target, r.lot,
                        r.lembar, _r(r.nilai_idr), _r(r.fee_beli_idr), _r(r.risiko_idr), r.risiko_pct,
                        _r(r.rr_bersih), r.instruksi_masuk, " ".join(r.ketergantungan), r.langkah_setelah,
                        " ".join(r.peringatan)])

    kandidat = d + [list(KOLOM_KANDIDAT)] + [[b[k] for k in KOLOM_KANDIDAT] for b in baris_kandidat(lap)]

    posisi = d + [["kode", "strategi", "saran", "alasan", "lembar", "harga_beli", "close", "stop_t+1", "target",
                   "hari_ke", "nilai_idr", "pnl_idr", "pnl_pct", "r", "perkiraan"]]
    for s in lap.posisi:
        posisi.append([s.kode, s.strategi, s.saran, s.alasan, s.lembar, s.harga_beli, s.close, s.stop_besok, s.target,
                       s.hari_ke, _r(s.nilai_idr), _r(s.pnl_idr), _r(s.pnl_pct), _r(s.r, 4), "≈" if s.perkiraan else ""])

    pasar = d + [["Status", status_pasar_ringkas(lap.kondisi)], [], ["syarat", "kondisi", "nilai", "ambang", "terpenuhi"]]
    if lap.kondisi is not None:
        pasar += [[s.aturan, s.uraian, s.nilai, s.ambang, s.tanda] for s in lap.kondisi.syarat]

    kualitas = d
    q = lap.kualitas
    if q is not None:
        kualitas = d + [["OK", q.status_ambil.get("OK", 0)], ["SEBAGIAN", q.status_ambil.get("SEBAGIAN", 0)],
                        ["GAGAL", q.status_ambil.get("GAGAL", 0)], ["data basi", q.basi], ["anomali", q.anomali],
                        ["sumber CSV manual", q.sumber_manual], [], ["kode", "alasan_dikeluarkan", "rincian"]]
        kualitas += [list(x) for x in q.dikeluarkan] + [[], ["Revisi data"]] + [[x] for x in q.revisi]

    catatan = d + [["Peringatan"]] + [[p] for p in lap.peringatan + lap.catatan]
    catatan += [[], ["Item belum diverifikasi"]] + [[p] for p in lap.belum_verifikasi]
    j = lap.jurnal
    catatan += [[], ["Jurnal kertas"]]
    if lap.jurnal_peringatan:
        catatan.append([f"Jurnal gagal: {lap.jurnal_peringatan}"])
    elif j is not None:
        catatan.append([f"{j.hari_dikejar} hari dikejar; {j.sinyal_susulan} sinyal susulan tidak pernah dilaporkan"])
        if j.ekuitas is not None:
            catatan.append([f"Ekuitas kertas {rp(j.ekuitas)} (periode {j.periode_id})"])
        for r in j.terbuka:
            catatan.append(["terbuka", *_baris_kertas(r)])
        for p in j.penyimpangan:
            catatan.append([f"Penyimpangan {p.strategi}: {p.catatan}"])
    catatan += [[], ["Keterbatasan"]] + [[x] for x in KETERBATASAN_SINGKAT]
    catatan += [[], [lap.catatan_skor], [DISCLAIMER]]
    return {"Ringkasan": ringkasan, "Rencana": rencana, "Kandidat": kandidat, "Posisi": posisi, "Pasar": pasar,
            "Kualitas Data": kualitas, "Catatan": catatan}


def bytes_xlsx(lap: LaporanHarian) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    isi = isi_sheet(lap)
    for nama in SHEET:
        ws = wb.create_sheet(nama)
        for baris in isi[nama]:
            ws.append(["" if x is None else x for x in baris] if baris else [])
        ws["A1"].font = Font(bold=True)
        ws.column_dimensions["A"].width = 18
    dibuat = lap.dibuat_pada.replace(tzinfo=None)
    wb.properties.created = dibuat
    wb.properties.modified = dibuat
    wb.properties.title = f"{lap.judul} — BUKAN SARAN INVESTASI"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
